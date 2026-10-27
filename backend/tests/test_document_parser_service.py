"""
Unit and Integration Tests for DocumentParserService.
Tests:
1. Azure DI success with polling workflow.
2. Missing Azure credentials (automatic local pdfplumber fallback).
3. Azure API failure (HTTP 500 / Network error) with graceful fallback to pdfplumber.
4. Azure polling timeout with graceful fallback to pdfplumber.
5. Empty Azure response followed by successful local pdfplumber extraction.
6. Multi-page PDFs (both Azure DI without 2-page artificial limit, and local pdfplumber).
7. Scanned / image-only PDFs (raises clear, user-friendly parsing error).
8. DOCX extraction (paragraphs, tables, word count) and rejection of legacy .doc files.
9. TXT extraction (UTF-8 safe, preserves meaningful newlines, rejects empty).
10. Missing optional dependencies (pdfplumber is None, python-docx is None).
11. Malformed files (corrupted PDF, corrupted DOCX, empty file, oversized file).
"""

import io
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from reportlab.pdfgen import canvas
from docx import Document

from services.document_parser_service import (
    DocumentParserService,
    MAX_FILE_SIZE_BYTES,
)


def create_sample_pdf(pages_text: list) -> bytes:
    """Helper to generate a real, valid multi-page PDF with selectable text."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages_text:
        c.drawString(100, 750, text)
        c.drawString(100, 700, "Skills: Python, FastAPI, Docker, Azure Cloud Services")
        c.drawString(100, 650, "Experience: 5 years of software engineering in distributed systems.")
        c.showPage()
    c.save()
    return buf.getvalue()


def create_sample_docx(paragraphs: list, table_data: list = None) -> bytes:
    """Helper to generate a valid DOCX with paragraphs and tables."""
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════
# 1. AZURE DOCUMENT INTELLIGENCE TESTS
# ══════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_azure_di_success():
    """Verifies successful asynchronous Azure DI submission and polling."""
    service = DocumentParserService(
        endpoint="https://mock-azure-di.cognitiveservices.azure.com",
        key="test-key-12345",
        timeout_seconds=5.0,
    )
    pdf_bytes = create_sample_pdf(["Page 1 Content", "Page 2 Content"])

    mock_submit_resp = MagicMock()
    mock_submit_resp.status_code = 202
    mock_submit_resp.headers = {"Operation-Location": "https://mock-azure-di/operations/op-123"}

    mock_poll_resp = MagicMock()
    mock_poll_resp.status_code = 200
    mock_poll_resp.json.return_value = {
        "status": "succeeded",
        "analyzeResult": {
            "content": "Jane Doe\nSoftware Engineer\nPython, Docker, Kubernetes\n5 years experience.",
            "pages": [{"pageNumber": 1}, {"pageNumber": 2}],
        },
    }

    mock_client = AsyncMock()
    mock_client.post = AsyncMock(return_value=mock_submit_resp)
    mock_client.get = AsyncMock(return_value=mock_poll_resp)

    with patch("httpx.AsyncClient") as mock_client_cls:
        mock_client_cls.return_value.__aenter__.return_value = mock_client
        res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "azure_di"
    assert res["page_count"] == 2
    assert "Jane Doe" in res["raw_text"]
    assert "Python, Docker" in res["raw_text"]
    assert res["metadata"]["provider"] == "azure_di"
    assert res["metadata"]["word_count"] > 5


@pytest.mark.asyncio
async def test_azure_di_no_two_page_limit():
    """Verifies that arbitrary two-page limit is removed and 5-page PDF succeeds with Azure DI."""
    service = DocumentParserService(
        endpoint="https://mock-azure-di.cognitiveservices.azure.com",
        key="test-key-12345",
        timeout_seconds=5.0,
    )
    pdf_bytes = create_sample_pdf([f"Page {i}" for i in range(1, 6)])

    mock_submit_resp = MagicMock()
    mock_submit_resp.status_code = 202
    mock_submit_resp.headers = {"Operation-Location": "https://mock-azure-di/operations/op-5page"}

    mock_poll_resp = MagicMock()
    mock_poll_resp.status_code = 200
    mock_poll_resp.json.return_value = {
        "status": "succeeded",
        "analyzeResult": {
            "content": "Long Resume Header\nPage 1 to Page 5 comprehensive technical achievements.",
            "pages": [{"pageNumber": i} for i in range(1, 6)],
        },
    }

    mock_client = AsyncMock()
    mock_client.post = AsyncMock(return_value=mock_submit_resp)
    mock_client.get = AsyncMock(return_value=mock_poll_resp)

    with patch("httpx.AsyncClient") as mock_client_cls:
        mock_client_cls.return_value.__aenter__.return_value = mock_client
        res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "azure_di"
    assert res["page_count"] == 5
    assert "Long Resume Header" in res["raw_text"]


@pytest.mark.asyncio
async def test_azure_di_missing_credentials_fallback():
    """Verifies that missing Azure credentials transparently use local pdfplumber without error."""
    service = DocumentParserService(endpoint="", key="")
    pdf_bytes = create_sample_pdf(["First Page Content"])

    res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "pdfplumber"
    assert res["page_count"] == 1
    assert "First Page Content" in res["raw_text"]
    assert res["metadata"]["fallback"] is False


@pytest.mark.asyncio
async def test_azure_di_api_failure_fallback():
    """Verifies that an Azure HTTP 500 error falls back to local pdfplumber."""
    service = DocumentParserService(
        endpoint="https://mock-azure-di.cognitiveservices.azure.com",
        key="test-key-12345",
    )
    pdf_bytes = create_sample_pdf(["Local Fallback Resume Page"])

    mock_submit_resp = MagicMock()
    mock_submit_resp.status_code = 500
    mock_submit_resp.text = "Internal Server Error"

    mock_client = AsyncMock()
    mock_client.post = AsyncMock(return_value=mock_submit_resp)

    with patch("httpx.AsyncClient") as mock_client_cls:
        mock_client_cls.return_value.__aenter__.return_value = mock_client
        res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "pdfplumber"
    assert res["page_count"] == 1
    assert "Local Fallback Resume Page" in res["raw_text"]
    assert res["metadata"]["fallback"] is True


@pytest.mark.asyncio
async def test_azure_di_timeout_fallback():
    """Verifies that an Azure polling timeout falls back to local pdfplumber."""
    service = DocumentParserService(
        endpoint="https://mock-azure-di.cognitiveservices.azure.com",
        key="test-key-12345",
        timeout_seconds=0.05,
    )
    pdf_bytes = create_sample_pdf(["Timeout Fallback Resume Page"])

    mock_submit_resp = MagicMock()
    mock_submit_resp.status_code = 202
    mock_submit_resp.headers = {"Operation-Location": "https://mock-azure-di/operations/op-timeout"}

    mock_poll_resp = MagicMock()
    mock_poll_resp.status_code = 200
    mock_poll_resp.json.return_value = {"status": "running"}

    mock_client = AsyncMock()
    mock_client.post = AsyncMock(return_value=mock_submit_resp)
    mock_client.get = AsyncMock(return_value=mock_poll_resp)

    with patch("httpx.AsyncClient") as mock_client_cls:
        mock_client_cls.return_value.__aenter__.return_value = mock_client
        res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "pdfplumber"
    assert "Timeout Fallback Resume Page" in res["raw_text"]
    assert res["metadata"]["fallback"] is True


@pytest.mark.asyncio
async def test_empty_azure_response_followed_by_local_extraction():
    """Verifies that an empty Azure DI text response triggers local pdfplumber extraction."""
    service = DocumentParserService(
        endpoint="https://mock-azure-di.cognitiveservices.azure.com",
        key="test-key-12345",
    )
    pdf_bytes = create_sample_pdf(["Real Selectable Local PDF Content"])

    mock_submit_resp = MagicMock()
    mock_submit_resp.status_code = 202
    mock_submit_resp.headers = {"Operation-Location": "https://mock-azure-di/operations/op-empty"}

    mock_poll_resp = MagicMock()
    mock_poll_resp.status_code = 200
    mock_poll_resp.json.return_value = {
        "status": "succeeded",
        "analyzeResult": {"content": "   ", "paragraphs": [], "pages": [{"pageNumber": 1}]},
    }

    mock_client = AsyncMock()
    mock_client.post = AsyncMock(return_value=mock_submit_resp)
    mock_client.get = AsyncMock(return_value=mock_poll_resp)

    with patch("httpx.AsyncClient") as mock_client_cls:
        mock_client_cls.return_value.__aenter__.return_value = mock_client
        res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "pdfplumber"
    assert "Real Selectable Local PDF Content" in res["raw_text"]
    assert res["metadata"]["fallback"] is True


# ══════════════════════════════════════════════════════════════════════════
# 2. LOCAL PDFPLUMBER MULTI-PAGE & SCANNED ERROR TESTS
# ══════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_multipage_pdf_pdfplumber():
    """Verifies multi-page local extraction via pdfplumber."""
    service = DocumentParserService(endpoint="", key="")
    pdf_bytes = create_sample_pdf(["Section A - Work History", "Section B - Education Details", "Section C - Projects"])

    res = await service.parse_document(pdf_bytes, "pdf")

    assert res["source"] == "pdfplumber"
    assert res["page_count"] == 3
    assert "Section A - Work History" in res["raw_text"]
    assert "Section B - Education Details" in res["raw_text"]
    assert "Section C - Projects" in res["raw_text"]


@pytest.mark.asyncio
async def test_scanned_image_only_pdf_raises_error():
    """Verifies that scanned / image-only PDFs with no selectable text raise a descriptive error."""
    service = DocumentParserService(endpoint="", key="")
    # Generate empty page PDF without text
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()
    empty_pdf = buf.getvalue()

    with pytest.raises(ValueError) as exc:
        await service.parse_document(empty_pdf, "pdf")

    assert "scanned image" in str(exc.value).lower() or "readable text" in str(exc.value).lower()


# ══════════════════════════════════════════════════════════════════════════
# 3. DOCX & TXT TESTS
# ══════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_docx_extraction_success():
    """Verifies extraction of paragraphs and tables from DOCX files."""
    service = DocumentParserService()
    paragraphs = [
        "Senior Software Engineer Resume",
        "Core skills: Python, React, FastAPI, MongoDB, PostgreSQL.",
    ]
    table = [
        ["Company", "Role", "Duration"],
        ["Acme Corp", "Backend Lead", "2021-2025"],
    ]
    docx_bytes = create_sample_docx(paragraphs, table)

    res = await service.parse_document(docx_bytes, "docx")

    assert res["source"] == "docx"
    assert res["page_count"] >= 1
    assert "Senior Software Engineer Resume" in res["raw_text"]
    assert "Acme Corp" in res["raw_text"]
    assert "Backend Lead" in res["raw_text"]
    assert res["metadata"]["provider"] == "docx"


@pytest.mark.asyncio
async def test_legacy_doc_rejection():
    """Verifies that legacy binary .doc format is rejected with explicit conversion guidance."""
    service = DocumentParserService()
    with pytest.raises(ValueError) as exc:
        await service.parse_document(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 binary doc", "doc")

    assert "Legacy .doc format is not supported" in str(exc.value)


@pytest.mark.asyncio
async def test_txt_extraction_success():
    """Verifies UTF-8 safe text extraction with meaningful newline preservation."""
    service = DocumentParserService()
    txt_content = "Line 1: Software Engineer\n\nLine 2: Skills in AI, ML & NLP\nLine 3: Étudiant diplômé 2026 🚀"

    res = await service.parse_document(txt_content.encode("utf-8"), "txt")

    assert res["source"] == "txt"
    assert res["page_count"] == 1
    assert "Line 1: Software Engineer" in res["raw_text"]
    assert "Line 2: Skills in AI" in res["raw_text"]
    assert "Étudiant diplômé 2026 🚀" in res["raw_text"]


@pytest.mark.asyncio
async def test_empty_txt_raises_error():
    """Verifies that empty TXT documents raise ValueError."""
    service = DocumentParserService()
    with pytest.raises(ValueError) as exc:
        await service.parse_document(b"   \n\t  \n  ", "txt")
    assert "empty" in str(exc.value).lower()


# ══════════════════════════════════════════════════════════════════════════
# 4. MISSING DEPENDENCIES & MALFORMED INPUT TESTS
# ══════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_missing_pdfplumber_raises_runtime_error(monkeypatch):
    """Verifies helpful error when pdfplumber is not installed and Azure is unavailable."""
    import services.document_parser_service as dps
    monkeypatch.setattr(dps, "pdfplumber", None)

    service = DocumentParserService(endpoint="", key="")
    pdf_bytes = create_sample_pdf(["Valid text"])

    with pytest.raises(RuntimeError) as exc:
        await service.parse_document(pdf_bytes, "pdf")
    assert "Neither Azure Document Intelligence nor pdfplumber is available" in str(exc.value)


@pytest.mark.asyncio
async def test_missing_docx_raises_runtime_error(monkeypatch):
    """Verifies helpful error when python-docx is not installed."""
    import services.document_parser_service as dps
    monkeypatch.setattr(dps, "Document", None)

    service = DocumentParserService()
    with pytest.raises(RuntimeError) as exc:
        await service.parse_document(b"dummy docx bytes", "docx")
    assert "python-docx library is not installed" in str(exc.value)


@pytest.mark.asyncio
async def test_empty_file_input_raises_error():
    """Verifies that a 0-byte file raises ValueError."""
    service = DocumentParserService()
    with pytest.raises(ValueError) as exc:
        await service.parse_document(b"", "pdf")
    assert "empty" in str(exc.value).lower()


@pytest.mark.asyncio
async def test_oversized_file_input_raises_error():
    """Verifies that files exceeding 50 MB are rejected."""
    service = DocumentParserService()
    oversized = b"X" * (MAX_FILE_SIZE_BYTES + 1024)
    with pytest.raises(ValueError) as exc:
        await service.parse_document(oversized, "pdf")
    assert "exceeds maximum supported limit" in str(exc.value)


@pytest.mark.asyncio
async def test_malformed_corrupted_pdf_raises_error():
    """Verifies corrupted PDF raises clear error message."""
    service = DocumentParserService(endpoint="", key="")
    corrupted_pdf = b"%PDF-1.4\ncorrupted gibberish content that is not a real PDF structure"

    with pytest.raises(ValueError) as exc:
        await service.parse_document(corrupted_pdf, "pdf")
    assert "Unable to parse PDF" in str(exc.value) or "corrupted" in str(exc.value).lower()


@pytest.mark.asyncio
async def test_malformed_corrupted_docx_raises_error():
    """Verifies corrupted DOCX raises clear error message."""
    service = DocumentParserService()
    corrupted_docx = b"PK\x03\x04corrupted zip stream that cannot be opened as docx"

    with pytest.raises(ValueError) as exc:
        await service.parse_document(corrupted_docx, "docx")
    assert "Unable to read DOCX file" in str(exc.value)
